"""
Модуль хранения данных в Excel-файле (openpyxl).

Все таблицы хранятся на отдельных листах одного файла shop.xlsx:
- customers (id, name, email, phone, discount)
- products  (id, name, price)
- orders    (id, customer_id, date)
- order_items (order_id, product_id, quantity)
"""

import os
import csv
import json
from openpyxl import Workbook, load_workbook
from models import Customer, Product, Order

FILE_PATH = os.path.join(os.path.dirname(__file__), "data", "shop.xlsx")

SHEETS = {
    "customers": ["id", "name", "email", "phone", "discount"],
    "products": ["id", "name", "price"],
    "orders": ["id", "customer_id", "date"],
    "order_items": ["order_id", "product_id", "quantity"],
}


class ExcelStorage:
    """Обёртка над Excel-файлом, эмулирующая БД."""

    def __init__(self, path: str = FILE_PATH):
        self.path = path
        os.makedirs(os.path.dirname(path), exist_ok=True)
        if not os.path.exists(path):
            self._create_empty()
        self.wb = load_workbook(path)

    # ---------- Внутренние утилиты ----------
    def _create_empty(self):
        """Создать новый Excel-файл со всеми листами."""
        wb = Workbook()
        # удалим дефолтный лист
        wb.remove(wb.active)
        for name, headers in SHEETS.items():
            ws = wb.create_sheet(name)
            ws.append(headers)
        wb.save(self.path)

    def _sheet_rows(self, sheet: str) -> list[dict]:
        """Прочитать строки листа как список словарей."""
        ws = self.wb[sheet]
        headers = [c.value for c in ws[1]]
        rows = []
        for row in ws.iter_rows(min_row=2, values_only=True):
            if row is None or all(v is None for v in row):
                continue
            rows.append(dict(zip(headers, row)))
        return rows

    def _append(self, sheet: str, values: list):
        """Добавить строку на лист."""
        self.wb[sheet].append(values)

    def _next_id(self, sheet: str) -> int:
        """Следующий id для листа."""
        rows = self._sheet_rows(sheet)
        if not rows:
            return 1
        return max(int(r["id"]) for r in rows) + 1

    def _save(self):
        """Сохранить файл."""
        self.wb.save(self.path)

    def reload(self):
        """Перечитать файл (если он был изменён извне)."""
        self.wb = load_workbook(self.path)

    # ---------- Клиенты ----------
    def add_customer(self, customer: Customer) -> int:
        """Добавить клиента, вернуть id."""
        cid = self._next_id("customers")
        self._append("customers", [
            cid, customer.name, customer.email, customer.phone,
            getattr(customer, "discount", 0.0),
        ])
        self._save()
        return cid

    def get_customers(self) -> list[Customer]:
        """Все клиенты."""
        result = []
        for r in self._sheet_rows("customers"):
            try:
                result.append(Customer(r["name"], r["email"], r["phone"], int(r["id"])))
            except ValueError:
                continue
        return result

    def delete_customer(self, customer_id: int) -> None:
        """Удалить клиента по id (пересобираем лист)."""
        ws = self.wb["customers"]
        headers = [c.value for c in ws[1]]
        new_rows = [headers]
        for row in ws.iter_rows(min_row=2, values_only=True):
            if row is None or all(v is None for v in row):
                continue
            if int(row[0]) != int(customer_id):
                new_rows.append(list(row))
        ws.delete_rows(1, ws.max_row)
        for r in new_rows:
            ws.append(r)
        self._save()

    def search_customers(self, query: str) -> list[Customer]:
        """Поиск по имени/email (лямбда + filter)."""
        q = query.lower().strip()
        return list(filter(lambda c: q in c.name.lower() or q in c.email.lower(),
                           self.get_customers()))

    # ---------- Товары ----------
    def add_product(self, product: Product) -> int:
        """Добавить товар."""
        pid = self._next_id("products")
        self._append("products", [pid, product.name, product.price])
        self._save()
        return pid

    def get_products(self) -> list[Product]:
        """Все товары."""
        result = []
        for r in self._sheet_rows("products"):
            result.append(Product(r["name"], float(r["price"]), int(r["id"])))
        return result

    # ---------- Заказы ----------
    def add_order(self, order: Order) -> int:
        """Добавить заказ и его позиции."""
        oid = self._next_id("orders")
        self._append("orders", [oid, order.customer_id, order.date])
        for pid, qty in order.items:
            self._append("order_items", [oid, pid, qty])
        self._save()
        return oid

    def get_orders(self) -> list[Order]:
        """Все заказы с позициями."""
        items_by_order: dict[int, list[tuple[int, int]]] = {}
        for r in self._sheet_rows("order_items"):
            items_by_order.setdefault(int(r["order_id"]), []).append(
                (int(r["product_id"]), int(r["quantity"]))
            )
        result = []
        for r in self._sheet_rows("orders"):
            oid = int(r["id"])
            result.append(Order(
                customer_id=int(r["customer_id"]),
                date=str(r["date"]),
                items=items_by_order.get(oid, []),
                order_id=oid,
            ))
        return result

    # ---------- CSV ----------
    def export_customers_csv(self, path: str) -> None:
        """Экспорт клиентов в CSV."""
        with open(path, "w", newline="", encoding="utf-8") as f:
            writer = csv.writer(f)
            writer.writerow(["id", "name", "email", "phone"])
            for c in self.get_customers():
                writer.writerow([c.id, c.name, c.email, c.phone])

    def import_customers_csv(self, path: str) -> int:
        """Импорт клиентов из CSV, вернуть число добавленных."""
        count = 0
        with open(path, "r", encoding="utf-8") as f:
            reader = csv.DictReader(f)
            for row in reader:
                try:
                    self.add_customer(Customer(row["name"], row["email"], row["phone"]))
                    count += 1
                except (ValueError, KeyError):
                    continue
        return count

    # ---------- JSON ----------
    def export_all_json(self, path: str) -> None:
        """Экспорт всех данных в JSON."""
        data = {
            "customers": [c.to_dict() for c in self.get_customers()],
            "products": [p.to_dict() for p in self.get_products()],
            "orders": [o.to_dict() for o in self.get_orders()],
        }
        with open(path, "w", encoding="utf-8") as f:
            json.dump(data, f, ensure_ascii=False, indent=2, default=str)

    def import_all_json(self, path: str) -> None:
        """Импорт данных из JSON."""
        with open(path, "r", encoding="utf-8") as f:
            data = json.load(f)
        for c in data.get("customers", []):
            try:
                self.add_customer(Customer(c["name"], c["email"], c["phone"]))
            except ValueError:
                pass
        for p in data.get("products", []):
            self.add_product(Product(p["name"], p["price"]))
        for o in data.get("orders", []):
            items = [tuple(i) for i in o.get("items", [])]
            self.add_order(Order(o["customer_id"], o["date"], items))